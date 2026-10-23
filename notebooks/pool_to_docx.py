import collections
import html
import json
import re
from pathlib import Path

import docx

ROOT = next(p for p in Path(__file__).resolve().parents if (p / "infra").is_dir())
DATASET_DIR = ROOT.parent / "datasets" / "annotated_dataset"
POOL = DATASET_DIR / "pool"
OUT = DATASET_DIR / "docx_red_flags"
TEMPLATE = ROOT / "infra/contracts_docx/Synthetic_OrvaneTalvera_Supply Agreement.docx"
RED_FLAG_CLASS = "e_23"

PAGE_MARKER = re.compile(
    r"^(?:-\s*)?(?:page\s+)?(?:\d{1,3}|[ivxlc]{1,6})(?:\s+of\s+\d{1,3})?(?:\s*-)?$", re.IGNORECASE
)
NUMBERISH = re.compile(r"^[\d.,$%()\-–\s]+$")  # un número o importe suelto, p. ej. una celda de tabla
CLAUSE_NUMBER = re.compile(r"^\d{1,2}(?:\.\d{1,2})*\.?$")
ITEM_MARKER = re.compile(r"^\((?:[a-z]{1,2}|[ivx]{1,5}|\d{1,2}|[A-Z])\)$")
COVER_LINES = 20  # un bloque de título más largo que esto es una tabla, no una portada
COVER_END = re.compile(r"^(?:TABLE OF CONTENTS|Table of Contents|INDEX|CONTENTS)\b")
NOT_A_TITLE = re.compile(r"^(?:ARTICLE|Article|SECTION|Section|EXHIBIT|Exhibit|SCHEDULE|Schedule|PAGE|Page)\b")
EXHIBIT = re.compile(r"^(?:Exhibit|EXHIBIT)\s+[\dA-Z][\w.\-()]*$")
ARTICLE = re.compile(r"^(?:ARTICLE|Article|SECTION|Section|PART|Part|CHAPTER|Chapter)\s+[\dIVXLC]+[A-Z]?\b")
SUB_ITEM = re.compile(
    r"^(?:\((?:[a-z]{1,2}|[ivx]{1,5}|\d{1,2}|[A-Z])\)|[a-z]\)|\d{1,2}\s?(?:st|nd|rd|th)\.-|[A-Z]\.\d+\.?)\s"
)
NUMBER = r"(?:(?:Section|SECTION|Article|ARTICLE)\s+)?(?:\d+(?:\.\d+)*\.?|[IVXLC]+\.|[A-Z]\.)"
ORDINAL = r"(?:[A-Z]+(?:TH|ST|ND|RD|ND)?|[A-Z][a-z]+(?:th|st|nd|rd))(?::|\.-)"
CAPTION = re.compile(rf"^((?:{NUMBER}|{ORDINAL})\s+)(.{{1,90}}?(?:\.-|\.|:))(?=\s|$)")
ABBREVIATION_END = re.compile(r"(?:\b[A-Z]\.){2,}$|\b(?:Inc|Ltd|Co|Corp|Bros|No|St|Mr|Ms|Dr)\.$")
SMALL_WORDS = {"a", "an", "and", "as", "at", "by", "for", "in", "of", "on", "or", "the", "to", "with", "upon", "under"}


def read_parts(path):
    s = Path(path).read_text(encoding="utf-8", errors="replace")
    return [
        (pid, html.unescape(re.sub(r"<[^>]+>", "", body)).strip())
        for pid, body in re.findall(r'<p id="([^"]+)">(.*?)</p>', s, re.S)
    ]


def mostly_upper(text):
    letters = [c for c in text if c.isalpha()]
    return len(letters) >= 3 and sum(c.isupper() for c in letters) / len(letters) >= 0.8


def is_heading(text):
    words = text.split()
    if ARTICLE.match(text) and len(words) <= 15:
        return True
    return mostly_upper(text) and len(words) <= 10 and not text.endswith((",", ";"))


def caption_split(text):
    """(número, epígrafe, resto) cuando la cláusula empieza con un título corto; si no, None."""
    m = CAPTION.match(text)
    if not m:
        return None
    lead, cap = m.group(1), m.group(2)
    if cap.count(",") >= 2 or ABBREVIATION_END.search(cap):  # es una frase, no un título
        return None
    words = re.findall(r"[A-Za-z][\w’'-]*", cap)
    if not words or len(words) > 10 or max(len(w) for w in words) < 3:  # «R. V.» es un nombre
        return None
    titled = [w for w in words if w.lower() not in SMALL_WORDS]
    title_case = titled and sum(w[0].isupper() for w in titled) / len(titled) >= 0.75
    if not (cap.endswith(".-") or title_case or mostly_upper(cap)):
        return None
    return lead, cap, text[m.end():]


def short_title(text):
    words = text.split()
    return (
        0 < len(words) <= 8
        and text[0].isupper()
        and not text.endswith((".", ",", ";"))
        and not NOT_A_TITLE.match(text)
        and not NUMBERISH.match(text)
    )


def clause_numbers(parts):
    """Líneas con solo un número de cláusula o un marcador «(i)» que una maquetación a dos
    columnas separó de su texto."""
    texts = [t for _, t in parts]
    dotted = [i for i, t in enumerate(texts) if CLAUSE_NUMBER.match(t) and "." in t.strip(".")]
    found = set()
    for i, t in enumerate(texts):
        if not CLAUSE_NUMBER.match(t) or i + 1 >= len(texts) or not texts[i + 1]:
            continue
        nxt = texts[i + 1]
        if NUMBERISH.match(nxt):
            continue
        if i in dotted or t.endswith(".") or (len(dotted) >= 3 and short_title(nxt)):
            found.add(i)
    for i, t in enumerate(texts[:-1]):
        if ITEM_MARKER.match(t) and texts[i + 1] and not NUMBERISH.match(texts[i + 1]):
            found.add(i)
    return found


def clean(parts):
    """Quita los restos de paginación y vuelve a unir las frases que partió un salto de página.

    Devuelve [(texto, [(id de línea, posición en el texto)])] y los ids que se quitaron.
    """
    texts = [t for _, t in parts]
    numbers = clause_numbers(parts)

    def isolated(i):
        # Un número de página va solo; varios números juntos son una tabla aplanada.
        window = range(max(0, i - 2), min(len(texts), i + 3))
        return not any(
            j != i and j not in numbers and texts[j] and NUMBERISH.match(texts[j]) for j in window
        )

    markers = {
        i for i, t in enumerate(texts) if PAGE_MARKER.match(t) and i not in numbers and isolated(i)
    }
    short = collections.defaultdict(list)
    for i, t in enumerate(texts):
        if t and len(t.split()) <= 8 and i not in markers:
            short[t].append(i)
    running = set()
    for t, idx in short.items():
        near = [i for i in idx if {i - 1, i + 1} & markers]
        if len(idx) >= 3 and len(near) >= len(idx) - 1:
            running.update(idx[1:])  # la primera suele ser el título de verdad

    out, dropped, after_break, pending = [], [], False, None
    for i, (pid, t) in enumerate(parts):
        if not t or i in markers or i in running:
            dropped.append(pid)
            after_break = after_break or i in markers or i in running
            continue
        if i in numbers:
            if pending:  # dos marcadores seguidos: el primero queda solo
                out.append((pending[1], [(pending[0], 0)]))
            pending = (pid, t)
            continue
        if t == ":" and out and not pending:  # «Term» / «:» / valor, de una tabla a dos columnas
            prev_text, prev_ids = out[-1]
            out[-1] = (prev_text + t, prev_ids + [(pid, len(prev_text))])
            continue
        ids = [(pid, 0)]
        if pending:
            ids = [(pending[0], 0), (pid, len(pending[1]) + 1)]
            t = pending[1] + " " + t
            pending = None
        if after_break and out and t[0].islower() and not re.search(r"[.:;!?)\"”]$", out[-1][0]):
            prev_text, prev_ids = out[-1]
            out[-1] = (prev_text + " " + t, prev_ids + [(p, o + len(prev_text) + 1) for p, o in ids])
        else:
            out.append((t, ids))
        after_break = False
    if pending:
        out.append((pending[1], [(pending[0], 0)]))
    return out, dropped


def style_of(text, idx, first_body, signing):
    if idx == 0 and EXHIBIT.match(text):
        return "Body Text.right", True
    if idx < first_body and len(text.split()) <= 14:
        return "Body Text.center", mostly_upper(text)
    numbered = re.match(r"^\d{1,2}(?:\.\d{1,2})*\.?\s+(.*)$", text)
    if numbered and short_title(numbered.group(1)):
        return "Body Text", True  # «5.2 Headlease»: un título numerado
    if not signing and is_heading(text):
        return "Body Text.center", True
    if SUB_ITEM.match(text):
        return "Body Text.item", False
    return "Body Text", False


def blank_document():
    """La plantilla sin su texto: conserva página, márgenes y estilos."""
    d = docx.Document(TEMPLATE)
    body = d.element.body
    for el in list(body):
        if el.tag.endswith("}sectPr"):
            continue
        body.remove(el)
    props = d.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.comments = ""
    props.subject = ""
    props.keywords = ""
    return d


def build(name, parts):
    paras, dropped = clean(parts)
    first_body = next(
        (i for i, (t, _) in enumerate(paras) if len(t.split()) > 25 or COVER_END.match(t)), len(paras)
    )
    first_body = first_body if first_body <= COVER_LINES else 1
    d = blank_document()
    d.core_properties.title = name
    mapping = {pid: None for pid in dropped}
    signing = False
    for i, (text, ids) in enumerate(paras):
        if re.match(r"^IN WITNESS WHEREOF", text, re.IGNORECASE):
            signing = True
        style, bold = style_of(text, i, first_body, signing)
        p = d.add_paragraph(style=style)
        split = caption_split(text) if style in ("Body Text", "Body Text.item") else None
        if split:
            lead, cap, rest = split
            p.add_run(lead + cap).bold = True
            p.add_run(rest)
        else:
            p.add_run(text).bold = bold or None
        for pid, off in ids:
            mapping[pid] = {"paragraph": i, "offset": off}
    missing = [pid for pid, _ in parts if pid not in mapping]
    assert not missing, f"{name}: líneas perdidas sin haberlas quitado: {missing[:5]}"
    return d, mapping


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    maps, summary = {}, []
    for ann in sorted(POOL.glob("*.ann.json")):
        entities = json.loads(ann.read_text())["entities"]
        n_flags = sum(1 for e in entities if e["classId"] == RED_FLAG_CLASS)
        if not n_flags:
            continue
        stem = ann.name[: -len(".ann.json")]
        name = stem.split("-", 1)[1]
        parts = read_parts(POOL / f"{stem}.plain.html")
        d, mapping = build(name, parts)
        d.save(OUT / f"{name}.docx")
        maps[name] = {"tagtog_id": stem, "parts": mapping}
        summary.append((name, n_flags, len(parts), sum(v is None for v in mapping.values())))
    (OUT / "paragraph_map.json").write_text(json.dumps(maps, indent=1))
    print(f"{len(summary)} contratos -> {OUT}")
    print("líneas descartadas (paginación):", sum(s[3] for s in summary), "de", sum(s[2] for s in summary))


if __name__ == "__main__":
    main()
